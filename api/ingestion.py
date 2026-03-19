"""
Document ingestion pipeline (v2.1/v2.2/v2.3).
Supports: pdf, txt, docx, csv.
Features: SHA-256 dedup, OCR (pdfplumber + Tesseract), OCR confidence scoring,
          token chunking, header-aware semantic chunking (v2.3, feature-flagged).
"""
import hashlib
import io
import os
from typing import Optional
from uuid import UUID

import pdfplumber
import pytesseract
import tiktoken
import openai
from fastapi import HTTPException

SUPPORTED_FILE_TYPES = {"pdf", "txt", "docx", "csv"}
EMBEDDING_MODEL = "text-embedding-3-small"

_tokenizer = tiktoken.get_encoding("cl100k_base")


def detect_type(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def count_tokens(text: str) -> int:
    return len(_tokenizer.encode(text))


def chunk_text(text: str, size: int = 512, overlap: int = 64) -> list[str]:
    """Split text into token-sized chunks with overlap."""
    tokens = _tokenizer.encode(text)
    chunks = []
    start = 0
    while start < len(tokens):
        end = min(start + size, len(tokens))
        chunks.append(_tokenizer.decode(tokens[start:end]))
        if end == len(tokens):
            break
        start += size - overlap
    return chunks


def _extract_pdf_headings(file_bytes: bytes) -> list[dict]:
    """Font-size heuristics to identify headings in PDF (v2.3 semantic chunking)."""
    headings = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_num, page in enumerate(pdf.pages):
            words = page.extract_words(extra_attrs=["size", "fontname"])
            if not words:
                continue
            sizes = [w.get("size", 12) for w in words if w.get("size")]
            if not sizes:
                continue
            median_size = sorted(sizes)[len(sizes) // 2]
            for word in words:
                if word.get("size", 12) > median_size * 1.3:
                    headings.append({"text": word["text"], "page": page_num})
    return headings


def _extract_docx_headings(file_bytes: bytes) -> list[dict]:
    """Extract headings from DOCX paragraph styles (v2.3 semantic chunking)."""
    import docx as python_docx
    doc = python_docx.Document(io.BytesIO(file_bytes))
    headings = []
    for para in doc.paragraphs:
        if para.style and para.style.name.startswith("Heading"):
            headings.append({"text": para.text, "level": para.style.name})
    return headings


def _split_at_headings(full_text: str, headings: list[dict]) -> list[dict]:
    """Split text at heading positions."""
    if not headings:
        return [{"heading": "", "text": full_text}]

    sections = []
    remaining = full_text
    current_heading = ""
    for h in headings:
        heading_text = h["text"]
        idx = remaining.find(heading_text)
        if idx == -1:
            continue
        before = remaining[:idx].strip()
        if before:
            sections.append({"heading": current_heading, "text": before})
        current_heading = heading_text
        remaining = remaining[idx + len(heading_text):].strip()

    if remaining:
        sections.append({"heading": current_heading, "text": remaining})

    return sections if sections else [{"heading": "", "text": full_text}]


def chunk_text_semantic(
    full_text: str,
    headings: list[dict],
    size: int,
    overlap: int,
) -> list[dict]:
    """
    Split at heading boundaries first, fall back to token chunking within oversized sections.
    Each chunk includes section_heading in metadata (v2.3).
    """
    sections = _split_at_headings(full_text, headings)
    chunks = []
    for section in sections:
        heading = section.get("heading", "")
        body = section["text"]
        if not body.strip():
            continue
        if count_tokens(body) <= size:
            chunks.append({
                "text": f"{heading}\n{body}".strip() if heading else body,
                "metadata": {"section_heading": heading},
            })
        else:
            sub_chunks = chunk_text(body, size=size, overlap=overlap)
            for i, sc in enumerate(sub_chunks):
                chunks.append({
                    "text": f"{heading}\n{sc}".strip() if heading and i == 0 else sc,
                    "metadata": {"section_heading": heading},
                })
    return chunks


async def embed(text: str) -> list[float]:
    client = openai.AsyncOpenAI(api_key=os.environ["OPENAI_API_KEY"])
    resp = await client.embeddings.create(model=EMBEDDING_MODEL, input=text)
    return resp.data[0].embedding


async def ingest_document(
    db,
    business_id: UUID,
    file_bytes: bytes,
    filename: str,
    semantic_chunking: bool = False,
) -> dict:
    """Full ingestion pipeline. Returns {reused, document_id, ocr_quality, chunk_count}."""
    file_type = detect_type(filename)
    if file_type not in SUPPORTED_FILE_TYPES:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type. Accepted: {', '.join(sorted(SUPPORTED_FILE_TYPES))}",
        )

    # Dedup check
    fingerprint = hashlib.sha256(file_bytes).hexdigest()
    existing = await db.fetchrow(
        "SELECT id, status FROM documents WHERE business_id=$1 AND content_fingerprint=$2 AND status='ready'",
        business_id,
        fingerprint,
    )
    if existing:
        return {"reused": True, "document_id": str(existing["id"])}

    doc_id = await db.fetchval(
        """
        INSERT INTO documents (business_id, filename, file_type, content_fingerprint, status)
        VALUES ($1, $2, $3, $4, 'processing') RETURNING id
        """,
        business_id, filename, file_type, fingerprint,
    )

    full_text = ""
    ocr_applied = False
    ocr_confidence: Optional[float] = None
    headings: list[dict] = []

    try:
        if file_type == "pdf":
            pages = []
            confidence_scores = []
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    if len(text.strip()) < 50:
                        img = page.to_image(resolution=300).original
                        ocr_data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
                        text = pytesseract.image_to_string(img)
                        page_confs = [
                            int(c) for c in ocr_data["conf"]
                            if str(c).lstrip("-").isdigit() and int(c) > 0
                        ]
                        if page_confs:
                            confidence_scores.append(sum(page_confs) / len(page_confs) / 100.0)
                        ocr_applied = True
                    pages.append(text)
            full_text = "\n\n".join(pages)
            ocr_confidence = (
                sum(confidence_scores) / len(confidence_scores) if confidence_scores else None
            )
            if semantic_chunking:
                headings = _extract_pdf_headings(file_bytes)

        elif file_type == "txt":
            full_text = file_bytes.decode("utf-8", errors="replace")

        elif file_type == "docx":
            import docx as python_docx
            doc = python_docx.Document(io.BytesIO(file_bytes))
            full_text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if semantic_chunking:
                headings = _extract_docx_headings(file_bytes)

        elif file_type == "csv":
            import csv
            reader = csv.DictReader(io.StringIO(file_bytes.decode("utf-8", errors="replace")))
            rows = [". ".join(f"{k}: {v}" for k, v in row.items() if v) for row in reader]
            full_text = "\n".join(rows)
            # CSV has no heading structure — falls back to token chunking

        # Fetch business chunking config
        config = await db.fetchrow(
            "SELECT chunk_size_tokens, chunk_overlap_tokens FROM businesses WHERE id=$1",
            business_id,
        )
        size = config["chunk_size_tokens"]
        overlap = config["chunk_overlap_tokens"]

        if semantic_chunking and headings:
            raw_chunks = chunk_text_semantic(full_text, headings, size, overlap)
        else:
            raw_chunks = [{"text": c, "metadata": {}} for c in chunk_text(full_text, size, overlap)]

        # Embed and store each chunk
        for i, chunk_obj in enumerate(raw_chunks):
            chunk_text_val = chunk_obj["text"]
            chunk_meta = chunk_obj.get("metadata", {})
            chunk_meta.update({
                "source": filename,
                "chunk_index": i,
                "ocr": ocr_applied,
                "file_type": file_type,
            })
            embedding = await embed(chunk_text_val)
            await db.execute(
                """
                INSERT INTO document_chunks
                  (business_id, document_id, chunk_index, content, embedding, metadata)
                VALUES ($1, $2, $3, $4, $5::vector, $6::jsonb)
                """,
                business_id, doc_id, i, chunk_text_val, embedding, chunk_meta,
            )

        await db.execute(
            """
            UPDATE documents
            SET status='ready', ocr_applied=$1, ocr_confidence=$2, chunk_count=$3
            WHERE id=$4
            """,
            ocr_applied, ocr_confidence, len(raw_chunks), doc_id,
        )

    except Exception as e:
        await db.execute(
            "UPDATE documents SET status='error', error_message=$1 WHERE id=$2",
            str(e), doc_id,
        )
        raise

    ocr_quality = None
    if ocr_confidence is not None:
        ocr_quality = "low" if ocr_confidence < 0.6 else "ok"

    return {
        "reused": False,
        "document_id": str(doc_id),
        "ocr_quality": ocr_quality,
        "chunk_count": len(raw_chunks),
    }
